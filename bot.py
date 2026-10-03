import logging
import json
import os
import shutil
import sqlite3
import tempfile
from contextlib import closing
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, ReplyKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)


(
    PAY_TYPE,
    PAY_PROJECT,
    PAY_AMOUNT,
    PAY_NAME,
    PAY_RECEIPT,
    ADMIN_SETTING_VALUE,
    ADMIN_PROJECT_TITLE,
    ADMIN_PROJECT_TARGET,
    ADMIN_PROJECT_DETAILS,
    ADMIN_EXPENSE_TITLE,
    ADMIN_EXPENSE_CATEGORY,
    ADMIN_EXPENSE_AMOUNT,
    ADMIN_EXPENSE_DETAILS,
    REVIEW_REJECT_REASON,
    ADMIN_INPUT,
    ADMIN_REASON,
) = range(16)

PAYMENT_TYPES = {
    "monthly": "اشتراك شهري",
    "support": "دعم شهري",
    "donation_general": "تبرع للرابطة",
    "donation_project": "تبرع لمشروع",
}

EXPENSE_CATEGORIES = ["فعاليات", "طباعة", "دعم طلاب", "مواصلات", "ضيافة", "أخرى"]

MAIN_MENU = ReplyKeyboardMarkup(
    [
        ["الرصيد الحالي", "الدفع"],
        ["المشاريع", "الكشف"],
        ["المنصرفات", "الإدارة"],
    ],
    resize_keyboard=True,
)


def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def money(value) -> str:
    try:
        return f"{float(value):,.2f} جنيه"
    except (TypeError, ValueError):
        return "0.00 جنيه"


def amount_from_text(text: str) -> float:
    amount = float(text.replace(",", "").strip())
    if amount <= 0:
        raise ValueError
    return amount


class Store:
    def __init__(self, path: str):
        self.path = path
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.init()

    def connect(self):
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        return conn

    def init(self):
        with closing(self.connect()) as conn:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS settings (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS projects (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    title TEXT NOT NULL,
                    target_amount REAL NOT NULL DEFAULT 0,
                    details TEXT NOT NULL DEFAULT '',
                    active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS transactions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    telegram_user_id INTEGER NOT NULL,
                    telegram_username TEXT,
                    payer_name TEXT NOT NULL,
                    payment_type TEXT NOT NULL,
                    project_id INTEGER,
                    amount REAL NOT NULL,
                    receipt_file_id TEXT NOT NULL,
                    receipt_kind TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    admin_note TEXT,
                    created_at TEXT NOT NULL,
                    reviewed_at TEXT,
                    FOREIGN KEY(project_id) REFERENCES projects(id)
                );

                CREATE TABLE IF NOT EXISTS expenses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    title TEXT NOT NULL,
                    category TEXT NOT NULL DEFAULT 'أخرى',
                    amount REAL NOT NULL,
                    details TEXT NOT NULL DEFAULT '',
                    admin_user_id INTEGER NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS audit_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    admin_id INTEGER NOT NULL,
                    action TEXT NOT NULL,
                    entity_type TEXT NOT NULL,
                    entity_id TEXT,
                    old_data TEXT,
                    new_data TEXT,
                    reason TEXT NOT NULL DEFAULT '',
                    timestamp TEXT NOT NULL
                );
                """
            )
            self.migrate(conn)
            defaults = {
                "monthly_subscription_amount": "0",
                "payment_details": "لم يتم إضافة تفاصيل الدفع بعد.",
            }
            for key, value in defaults.items():
                conn.execute(
                    "INSERT OR IGNORE INTO settings(key, value) VALUES(?, ?)",
                    (key, value),
                )
            conn.commit()

    def migrate(self, conn):
        expense_columns = {
            row["name"]
            for row in conn.execute("PRAGMA table_info(expenses)").fetchall()
        }
        if "category" not in expense_columns:
            conn.execute("ALTER TABLE expenses ADD COLUMN category TEXT NOT NULL DEFAULT 'أخرى'")
        for table in ("transactions", "expenses"):
            columns = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}
            if "void_reason" not in columns:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN void_reason TEXT")
            if "modified_at" not in columns:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN modified_at TEXT")

    def audit(self, admin_id, action, entity_type, entity_id=None, old_data=None, new_data=None, reason=""):
        def safe_data(value):
            if isinstance(value, dict):
                return {k: safe_data(v) for k, v in value.items() if k.lower() not in {"bot_token", "token", "receipt_file_id"}}
            if isinstance(value, (list, tuple)):
                return [safe_data(v) for v in value]
            return value
        with closing(self.connect()) as conn:
            conn.execute("INSERT INTO audit_log(admin_id,action,entity_type,entity_id,old_data,new_data,reason,timestamp) VALUES(?,?,?,?,?,?,?,?)",
                         (admin_id, action, entity_type, str(entity_id) if entity_id is not None else None,
                          json.dumps(safe_data(old_data), ensure_ascii=False, default=str) if old_data is not None else None,
                          json.dumps(safe_data(new_data), ensure_ascii=False, default=str) if new_data is not None else None,
                          reason, now_iso()))
            conn.commit()

    def search(self, query, limit=10):
        term=query.strip().lstrip("@")
        needle = f"%{term}%"
        with closing(self.connect()) as conn:
            tx = conn.execute("SELECT t.*,p.title AS project_title FROM transactions t LEFT JOIN projects p ON p.id=t.project_id WHERE CAST(t.id AS TEXT) LIKE ? OR t.payer_name LIKE ? OR COALESCE(t.telegram_username,'') LIKE ? OR CAST(t.amount AS TEXT) LIKE ? OR t.status LIKE ? OR COALESCE(p.title,'') LIKE ? ORDER BY t.created_at DESC LIMIT ?", (needle,needle,needle,needle,needle,needle,limit)).fetchall()
            ex = conn.execute("SELECT * FROM expenses WHERE CAST(id AS TEXT) LIKE ? OR title LIKE ? OR details LIKE ? OR CAST(amount AS TEXT) LIKE ? ORDER BY created_at DESC LIMIT ?", (needle,needle,needle,needle,limit)).fetchall()
            return tx, ex

    def admin_transactions(self, limit=15):
        with closing(self.connect()) as conn:
            return conn.execute("SELECT t.*,p.title AS project_title FROM transactions t LEFT JOIN projects p ON p.id=t.project_id ORDER BY t.created_at DESC LIMIT ?", (limit,)).fetchall()

    def admin_expenses(self, limit=15):
        return self.expenses(limit, include_voided=True)

    def edit_transaction(self, tx_id, field, value, admin_id, reason):
        if field not in {"payer_name", "amount", "payment_type", "project_id"}:
            return False
        if field == "amount" and float(value) <= 0: return False
        if field == "payment_type" and value not in PAYMENT_TYPES: return False
        with closing(self.connect()) as conn:
            row = conn.execute("SELECT * FROM transactions WHERE id=?", (tx_id,)).fetchone()
            if not row or row["status"] == "voided": return False
            if field == "project_id" and value is not None:
                project=conn.execute("SELECT active FROM projects WHERE id=?", (value,)).fetchone()
                if not project or not project["active"]: return False
            old = dict(row); conn.execute(f"UPDATE transactions SET {field}=?, modified_at=? WHERE id=?", (value,now_iso(),tx_id)); new = dict(conn.execute("SELECT * FROM transactions WHERE id=?", (tx_id,)).fetchone()); conn.commit()
        self.audit(admin_id,"edit","transaction",tx_id,old,new,reason); return True

    def void_transaction(self, tx_id, admin_id, reason):
        with closing(self.connect()) as conn:
            row=conn.execute("SELECT * FROM transactions WHERE id=?",(tx_id,)).fetchone()
            if not row or row["status"]=="voided": return False
            old=dict(row); conn.execute("UPDATE transactions SET status='voided',void_reason=?,admin_note=?,reviewed_at=?,modified_at=? WHERE id=?",(reason,reason,now_iso(),now_iso(),tx_id)); new=dict(conn.execute("SELECT * FROM transactions WHERE id=?",(tx_id,)).fetchone()); conn.commit()
        self.audit(admin_id,"void","transaction",tx_id,old,new,reason); return True

    def edit_expense(self, expense_id, field, value, admin_id, reason):
        if field not in {"title", "category", "amount", "details"}: return False
        if field == "amount" and float(value) <= 0: return False
        with closing(self.connect()) as conn:
            row=conn.execute("SELECT * FROM expenses WHERE id=?",(expense_id,)).fetchone()
            if not row or row["void_reason"]: return False
            old=dict(row); conn.execute(f"UPDATE expenses SET {field}=?,modified_at=? WHERE id=?",(value,now_iso(),expense_id)); new=dict(conn.execute("SELECT * FROM expenses WHERE id=?",(expense_id,)).fetchone()); conn.commit()
        self.audit(admin_id,"edit","expense",expense_id,old,new,reason); return True

    def void_expense(self, expense_id, admin_id, reason):
        with closing(self.connect()) as conn:
            row=conn.execute("SELECT * FROM expenses WHERE id=?",(expense_id,)).fetchone()
            if not row or row["void_reason"]: return False
            old=dict(row); conn.execute("UPDATE expenses SET void_reason=?,modified_at=? WHERE id=?",(reason,now_iso(),expense_id)); new=dict(conn.execute("SELECT * FROM expenses WHERE id=?",(expense_id,)).fetchone()); conn.commit()
        self.audit(admin_id,"void","expense",expense_id,old,new,reason); return True

    def project_collected(self, project_id):
        with closing(self.connect()) as conn:
            return float(conn.execute("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE project_id=? AND status='approved'", (project_id,)).fetchone()[0])

    def reset_counts(self):
        with closing(self.connect()) as conn:
            return {"transactions":conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0], "expenses":conn.execute("SELECT COUNT(*) FROM expenses").fetchone()[0], "projects":conn.execute("SELECT COUNT(*) FROM projects").fetchone()[0]}

    def update_project(self, project_id, field, value, admin_id, reason):
        if field not in {"title", "target_amount", "details", "active"}: return False
        with closing(self.connect()) as conn:
            row=conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()
            if not row: return False
            old=dict(row); conn.execute(f"UPDATE projects SET {field}=? WHERE id=?",(value,project_id)); new=dict(conn.execute("SELECT * FROM projects WHERE id=?",(project_id,)).fetchone()); conn.commit()
        self.audit(admin_id,"update" if field!="active" else ("activate" if value else "deactivate"),"project",project_id,old,new,reason); return True

    def create_project_admin(self, title, target_amount, details, admin_id, reason="created by admin"):
        project_id=self.add_project(title,target_amount,details); self.audit(admin_id,"create","project",project_id,None,dict(self.project(project_id)),reason); return project_id

    def setting(self, key: str) -> str:
        with closing(self.connect()) as conn:
            row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
            return row["value"] if row else ""

    def set_setting(self, key: str, value: str):
        with closing(self.connect()) as conn:
            conn.execute(
                "INSERT INTO settings(key, value) VALUES(?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (key, value),
            )
            conn.commit()

    def add_project(self, title: str, target_amount: float, details: str) -> int:
        with closing(self.connect()) as conn:
            cursor = conn.execute(
                "INSERT INTO projects(title, target_amount, details, created_at) VALUES(?, ?, ?, ?)",
                (title, target_amount, details, now_iso()),
            )
            conn.commit()
            return int(cursor.lastrowid)

    def projects(self, active_only=True):
        where = "WHERE p.active = 1" if active_only else ""
        with closing(self.connect()) as conn:
            return conn.execute(
                f"""
                SELECT p.*,
                       COALESCE(SUM(CASE WHEN t.status = 'approved' THEN t.amount ELSE 0 END), 0) AS collected
                FROM projects p
                LEFT JOIN transactions t ON t.project_id = p.id
                {where}
                GROUP BY p.id
                ORDER BY p.created_at DESC
                """
            ).fetchall()

    def project(self, project_id: int):
        with closing(self.connect()) as conn:
            return conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()

    def add_transaction(self, data: dict) -> int:
        with closing(self.connect()) as conn:
            cursor = conn.execute(
                """
                INSERT INTO transactions (
                    telegram_user_id, telegram_username, payer_name, payment_type, project_id,
                    amount, receipt_file_id, receipt_kind, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    data["telegram_user_id"],
                    data.get("telegram_username"),
                    data["payer_name"],
                    data["payment_type"],
                    data.get("project_id"),
                    data["amount"],
                    data["receipt_file_id"],
                    data["receipt_kind"],
                    now_iso(),
                ),
            )
            conn.commit()
            return int(cursor.lastrowid)

    def update_transaction_status(self, tx_id: int, status: str, admin_note: str = ""):
        with closing(self.connect()) as conn:
            conn.execute(
                """
                UPDATE transactions
                SET status = ?, admin_note = ?, reviewed_at = ?
                WHERE id = ? AND status = 'pending'
                """,
                (status, admin_note, now_iso(), tx_id),
            )
            conn.commit()

    def transaction(self, tx_id: int):
        with closing(self.connect()) as conn:
            return conn.execute(
                """
                SELECT t.*, p.title AS project_title
                FROM transactions t
                LEFT JOIN projects p ON p.id = t.project_id
                WHERE t.id = ?
                """,
                (tx_id,),
            ).fetchone()

    def pending_transactions(self):
        with closing(self.connect()) as conn:
            return conn.execute(
                """
                SELECT t.*, p.title AS project_title
                FROM transactions t
                LEFT JOIN projects p ON p.id = t.project_id
                WHERE t.status = 'pending'
                ORDER BY t.created_at ASC
                LIMIT 20
                """
            ).fetchall()

    def approved_transactions(self, limit=15):
        with closing(self.connect()) as conn:
            return conn.execute(
                """
                SELECT t.*, p.title AS project_title
                FROM transactions t
                LEFT JOIN projects p ON p.id = t.project_id
                WHERE t.status = 'approved'
                ORDER BY t.reviewed_at DESC, t.created_at DESC
                LIMIT ?
                """,
                (limit,),
            ).fetchall()

    def add_expense(self, title: str, category: str, amount: float, details: str, admin_user_id: int) -> int:
        with closing(self.connect()) as conn:
            cursor = conn.execute(
                "INSERT INTO expenses(title, category, amount, details, admin_user_id, created_at) VALUES(?, ?, ?, ?, ?, ?)",
                (title, category, amount, details, admin_user_id, now_iso()),
            )
            conn.commit()
            return int(cursor.lastrowid)

    def expenses(self, limit=15, include_voided=False):
        with closing(self.connect()) as conn:
            where = "" if include_voided else "WHERE void_reason IS NULL"
            return conn.execute(
                f"SELECT * FROM expenses {where} ORDER BY created_at DESC LIMIT ?",
                (limit,),
            ).fetchall()

    def expense_totals_by_category(self):
        with closing(self.connect()) as conn:
            return conn.execute(
                """
                SELECT category, COALESCE(SUM(amount), 0) AS total, COUNT(*) AS count
                FROM expenses
                WHERE void_reason IS NULL
                GROUP BY category
                ORDER BY total DESC
                """
            ).fetchall()

    def income_totals_by_type(self):
        with closing(self.connect()) as conn:
            return conn.execute(
                """
                SELECT payment_type, COALESCE(SUM(amount), 0) AS total, COUNT(*) AS count
                FROM transactions
                WHERE status = 'approved'
                GROUP BY payment_type
                ORDER BY total DESC
                """
            ).fetchall()

    def totals(self):
        with closing(self.connect()) as conn:
            income = conn.execute(
                "SELECT COALESCE(SUM(amount), 0) AS total FROM transactions WHERE status = 'approved'"
            ).fetchone()["total"]
            expenses = conn.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM expenses WHERE void_reason IS NULL").fetchone()[
                "total"
            ]
            pending = conn.execute("SELECT COUNT(*) AS count FROM transactions WHERE status = 'pending'").fetchone()[
                "count"
            ]
            return {
                "income": float(income),
                "expenses": float(expenses),
                "balance": float(income) - float(expenses),
                "pending": int(pending),
            }

    def monthly_totals(self, month_prefix: str):
        with closing(self.connect()) as conn:
            income = conn.execute(
                """
                SELECT COALESCE(SUM(amount), 0) AS total, COUNT(*) AS count
                FROM transactions
                WHERE status = 'approved' AND substr(COALESCE(reviewed_at, created_at), 1, 7) = ?
                """,
                (month_prefix,),
            ).fetchone()
            expenses = conn.execute(
                """
                SELECT COALESCE(SUM(amount), 0) AS total, COUNT(*) AS count
                FROM expenses
                WHERE void_reason IS NULL
                  AND substr(created_at, 1, 7) = ?
                """,
                (month_prefix,),
            ).fetchone()
            by_category = conn.execute(
                """
                SELECT category, COALESCE(SUM(amount), 0) AS total
                FROM expenses
                WHERE void_reason IS NULL AND substr(created_at, 1, 7) = ?
                GROUP BY category
                ORDER BY total DESC
                """,
                (month_prefix,),
            ).fetchall()
            by_project = conn.execute(
                """
                SELECT COALESCE(p.title, 'تبرعات عامة') AS title, COALESCE(SUM(t.amount), 0) AS total
                FROM transactions t
                LEFT JOIN projects p ON p.id = t.project_id
                WHERE t.status = 'approved'
                  AND t.payment_type IN ('donation_general', 'donation_project')
                  AND substr(COALESCE(t.reviewed_at, t.created_at), 1, 7) = ?
                GROUP BY COALESCE(p.title, 'تبرعات عامة')
                ORDER BY total DESC
                LIMIT 5
                """,
                (month_prefix,),
            ).fetchall()
            return {
                "income": float(income["total"]),
                "income_count": int(income["count"]),
                "expenses": float(expenses["total"]),
                "expense_count": int(expenses["count"]),
                "balance": float(income["total"]) - float(expenses["total"]),
                "by_category": by_category,
                "by_project": by_project,
            }

    def export_data(self):
        with closing(self.connect()) as conn:
            transactions = conn.execute(
                """
                SELECT t.*, p.title AS project_title
                FROM transactions t
                LEFT JOIN projects p ON p.id = t.project_id
                ORDER BY t.created_at DESC
                """
            ).fetchall()
            expenses = conn.execute("SELECT * FROM expenses ORDER BY created_at DESC").fetchall()
            projects = conn.execute(
                """
                SELECT p.*,
                       COALESCE(SUM(CASE WHEN t.status = 'approved' THEN t.amount ELSE 0 END), 0) AS collected
                FROM projects p
                LEFT JOIN transactions t ON t.project_id = p.id
                GROUP BY p.id
                ORDER BY p.created_at DESC
                """
            ).fetchall()
            return transactions, expenses, projects


def parse_admins() -> set[int]:
    admins = set()
    for value in os.getenv("ADMIN_IDS", "").split(","):
        value = value.strip()
        if value:
            admins.add(int(value))
    return admins


def is_admin(user_id: int | None) -> bool:
    return bool(user_id and user_id in ADMINS)


def admin_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [InlineKeyboardButton('المعاملات', callback_data='admin:transactions'), InlineKeyboardButton('المنصرفات', callback_data='admin:expenses')],
        [InlineKeyboardButton('المشاريع', callback_data='admin:projects'), InlineKeyboardButton('بحث', callback_data='admin:search')],
        [InlineKeyboardButton('نسخة احتياطية الآن', callback_data='admin:backup'), InlineKeyboardButton('إعادة ضبط آمنة', callback_data='admin:reset')],
        [InlineKeyboardButton('المدفوعات المعلقة', callback_data='admin:pending')],
        [InlineKeyboardButton('الإعدادات المالية', callback_data='admin:settings')],
        [InlineKeyboardButton('إضافة مشروع', callback_data='admin:add_project'), InlineKeyboardButton('إضافة منصرف', callback_data='admin:add_expense')],
        [InlineKeyboardButton('تقرير سريع', callback_data='admin:report'), InlineKeyboardButton('تقرير الشهر', callback_data='admin:monthly_report')],
        [InlineKeyboardButton('تصدير Excel', callback_data='admin:export_excel')],
    ])

def expense_category_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        [[InlineKeyboardButton(category, callback_data=f"expense_category:{category}")]
         for category in EXPENSE_CATEGORIES]
    )


def tx_summary(tx) -> str:
    project = f"\nالمشروع: {tx['project_title']}" if tx["project_title"] else ""
    username = f"@{tx['telegram_username']}" if tx["telegram_username"] else "بدون اسم مستخدم"
    return (
        f"طلب دفع رقم #{tx['id']}\n"
        f"الاسم: {tx['payer_name']}\n"
        f"الحساب: {username}\n"
        f"النوع: {PAYMENT_TYPES.get(tx['payment_type'], tx['payment_type'])}{project}\n"
        f"المبلغ: {money(tx['amount'])}\n"
        f"الحالة: {tx['status']}\n"
        f"التاريخ: {tx['created_at']}"
    )


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.effective_message.reply_text(
        "أهلا بيك في بوت المكتب المالي لرابطة الطلاب السودانيين بجامعة بورسعيد.\n"
        "اختار من القائمة:",
        reply_markup=MAIN_MENU,
    )


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.effective_message.reply_text("تم الإلغاء.", reply_markup=MAIN_MENU)
    return ConversationHandler.END


async def main_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    if text == "الرصيد الحالي":
        await show_public_balance(update, context)
    elif text == "المشاريع":
        await show_projects(update, context)
    elif text == "الكشف":
        await show_income_report(update, context)
    elif text == "المنصرفات":
        await show_expenses(update, context)
    elif text == "الإدارة":
        await show_admin_message(update, context)
    else:
        await update.message.reply_text("اختار من القائمة أو اكتب /start.", reply_markup=MAIN_MENU)


async def payment_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    keyboard = [
        [InlineKeyboardButton("اشتراك شهري", callback_data="paytype:monthly")],
        [InlineKeyboardButton("دعم شهري", callback_data="paytype:support")],
        [InlineKeyboardButton("تبرع للرابطة", callback_data="paytype:donation_general")],
        [InlineKeyboardButton("تبرع لمشروع", callback_data="paytype:donation_project")],
    ]
    await update.effective_message.reply_text("اختار نوع الدفع:", reply_markup=InlineKeyboardMarkup(keyboard))
    return PAY_TYPE


async def pay_type_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    payment_type = query.data.split(":", 1)[1]
    context.user_data["payment_type"] = payment_type

    if payment_type == "donation_project":
        projects = STORE.projects()
        if not projects:
            await query.edit_message_text("لا توجد مشاريع مضافة حاليا. يمكنك التبرع للرابطة مباشرة.")
            return ConversationHandler.END
        keyboard = [[InlineKeyboardButton(p["title"], callback_data=f"payproject:{p['id']}")] for p in projects]
        await query.edit_message_text("اختار المشروع:", reply_markup=InlineKeyboardMarkup(keyboard))
        return PAY_PROJECT

    await query.edit_message_text(payment_instructions(payment_type))
    return PAY_AMOUNT


async def pay_project_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    project_id = int(query.data.split(":", 1)[1])
    project = STORE.project(project_id)
    if not project or not project["active"]:
        await query.edit_message_text("المشروع غير متاح أو متوقف عن استقبال التبرعات.")
        return ConversationHandler.END
    context.user_data["project_id"] = project_id
    await query.edit_message_text(payment_instructions("donation_project", project["title"]))
    return PAY_AMOUNT


def payment_instructions(payment_type: str, project_title: str | None = None) -> str:
    monthly = STORE.setting("monthly_subscription_amount")
    details = STORE.setting("payment_details")
    target = f"\nالمشروع: {project_title}" if project_title else ""
    subscription = f"\nقيمة الاشتراك الشهري المحددة: {money(monthly)}" if payment_type == "monthly" else ""
    return (
        f"نوع الدفع: {PAYMENT_TYPES[payment_type]}{target}{subscription}\n\n"
        f"تفاصيل الدفع:\n{details}\n\n"
        "بعد الدفع اكتب المبلغ المدفوع بالأرقام."
    )


async def pay_amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        context.user_data["amount"] = amount_from_text(update.message.text)
    except ValueError:
        await update.message.reply_text("اكتب المبلغ بالأرقام فقط، مثال: 500")
        return PAY_AMOUNT
    await update.message.reply_text("اكتب اسمك الكامل كما تريد ظهوره في الكشف:")
    return PAY_NAME


async def pay_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = update.message.text.strip()
    if len(name) < 3:
        await update.message.reply_text("اكتب الاسم الكامل.")
        return PAY_NAME
    context.user_data["payer_name"] = name
    await update.message.reply_text("ارفق إشعار الدفع كصورة أو ملف PDF/صورة.")
    return PAY_RECEIPT


async def pay_receipt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    receipt_file_id = None
    receipt_kind = None
    if update.message.photo:
        receipt_file_id = update.message.photo[-1].file_id
        receipt_kind = "photo"
    elif update.message.document:
        receipt_file_id = update.message.document.file_id
        receipt_kind = "document"

    if not receipt_file_id:
        await update.message.reply_text("أرسل الإشعار كصورة أو ملف.")
        return PAY_RECEIPT

    user = update.effective_user
    tx_id = STORE.add_transaction(
        {
            **context.user_data,
            "telegram_user_id": user.id,
            "telegram_username": user.username,
            "receipt_file_id": receipt_file_id,
            "receipt_kind": receipt_kind,
        }
    )
    tx = STORE.transaction(tx_id)
    await notify_admins(context, tx)

    context.user_data.clear()
    await update.message.reply_text(
        f"تم إرسال طلب الدفع رقم #{tx_id} للإدارة. ستصلك رسالة بعد الاعتماد أو الرفض.",
        reply_markup=MAIN_MENU,
    )
    return ConversationHandler.END


async def notify_admins(context: ContextTypes.DEFAULT_TYPE, tx):
    keyboard = InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("اعتماد", callback_data=f"review:approve:{tx['id']}"),
                InlineKeyboardButton("رفض", callback_data=f"review:reject:{tx['id']}"),
            ]
        ]
    )
    for admin_id in ADMINS:
        try:
            await context.bot.send_message(admin_id, tx_summary(tx), reply_markup=keyboard)
            if tx["receipt_kind"] == "photo":
                await context.bot.send_photo(admin_id, tx["receipt_file_id"])
            else:
                await context.bot.send_document(admin_id, tx["receipt_file_id"])
        except Exception:
            logging.exception("Failed to notify admin %s", admin_id)


async def review_transaction(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    if not is_admin(query.from_user.id):
        await query.edit_message_text("هذه العملية مخصصة للإدارة فقط.")
        return ConversationHandler.END

    _, action, raw_tx_id = query.data.split(":")
    tx_id = int(raw_tx_id)
    tx = STORE.transaction(tx_id)
    if not tx or tx["status"] != "pending":
        await query.edit_message_text("هذا الطلب تمت مراجعته أو غير موجود.")
        return ConversationHandler.END

    if action == "reject":
        context.user_data.clear()
        context.user_data["reject_tx_id"] = tx_id
        await query.edit_message_text(
            f"اكتب سبب رفض طلب رقم #{tx_id}.\n"
            "مثال: الإشعار غير واضح أو المبلغ غير مطابق.\n"
            "للإلغاء اكتب /cancel"
        )
        return REVIEW_REJECT_REASON

    STORE.update_transaction_status(tx_id, "approved")
    STORE.audit(query.from_user.id, "approve", "transaction", tx_id, dict(tx), dict(STORE.transaction(tx_id)), "approved payment")
    tx = STORE.transaction(tx_id)
    status_ar = "تم اعتماد الدفع"

    await query.edit_message_text(f"{status_ar}.\n\n{tx_summary(tx)}")
    try:
        await context.bot.send_message(
            tx["telegram_user_id"],
            f"{status_ar} لطلب رقم #{tx_id}.\nالمبلغ: {money(tx['amount'])}",
        )
    except Exception:
        logging.exception("Failed to notify payer %s", tx["telegram_user_id"])
    return ConversationHandler.END


async def reject_reason(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END

    tx_id = context.user_data.get("reject_tx_id")
    reason = update.message.text.strip()
    if not tx_id:
        await update.message.reply_text("انتهت جلسة الرفض. افتح طلب الدفع مرة أخرى.", reply_markup=MAIN_MENU)
        return ConversationHandler.END
    if len(reason) < 3:
        await update.message.reply_text("اكتب سبب أوضح للرفض.")
        return REVIEW_REJECT_REASON

    tx = STORE.transaction(int(tx_id))
    if not tx or tx["status"] != "pending":
        await update.message.reply_text("هذا الطلب تمت مراجعته أو غير موجود.", reply_markup=MAIN_MENU)
        context.user_data.clear()
        return ConversationHandler.END

    STORE.update_transaction_status(int(tx_id), "rejected", reason)
    STORE.audit(update.effective_user.id, "reject", "transaction", int(tx_id), dict(tx), dict(STORE.transaction(int(tx_id))), reason)
    tx = STORE.transaction(int(tx_id))
    context.user_data.clear()

    await update.message.reply_text(f"تم رفض الدفع.\n\n{tx_summary(tx)}\nسبب الرفض: {reason}", reply_markup=MAIN_MENU)
    try:
        await context.bot.send_message(
            tx["telegram_user_id"],
            f"تم رفض الدفع لطلب رقم #{tx_id}.\n"
            f"المبلغ: {money(tx['amount'])}\n"
            f"سبب الرفض: {reason}",
        )
    except Exception:
        logging.exception("Failed to notify payer %s", tx["telegram_user_id"])
    return ConversationHandler.END


async def show_public_balance(update: Update, context: ContextTypes.DEFAULT_TYPE):
    totals = STORE.totals()
    month_prefix = datetime.now().strftime("%Y-%m")
    monthly = STORE.monthly_totals(month_prefix)
    income_by_type = STORE.income_totals_by_type()
    expense_by_category = STORE.expense_totals_by_category()

    lines = [
        "ملخص مالي عام",
        f"الرصيد الحالي: {money(totals['balance'])}",
        f"إجمالي الإيرادات المعتمدة: {money(totals['income'])}",
        f"إجمالي المنصرفات: {money(totals['expenses'])}",
        "",
        f"ملخص شهر {month_prefix}:",
        f"إيرادات الشهر: {money(monthly['income'])}",
        f"منصرفات الشهر: {money(monthly['expenses'])}",
        f"صافي الشهر: {money(monthly['balance'])}",
    ]

    lines.append("\nالإيرادات حسب النوع:")
    if income_by_type:
        for row in income_by_type:
            label = PAYMENT_TYPES.get(row["payment_type"], row["payment_type"])
            lines.append(f"- {label}: {money(row['total'])} ({row['count']} عملية)")
    else:
        lines.append("- لا توجد إيرادات معتمدة حتى الآن.")

    lines.append("\nالمنصرفات حسب التصنيف:")
    if expense_by_category:
        for row in expense_by_category:
            lines.append(f"- {row['category']}: {money(row['total'])} ({row['count']} عملية)")
    else:
        lines.append("- لا توجد منصرفات مسجلة حتى الآن.")

    await update.message.reply_text("\n".join(lines))


async def show_projects(update: Update, context: ContextTypes.DEFAULT_TYPE):
    projects = STORE.projects()
    if not projects:
        await update.message.reply_text("لا توجد مشاريع مضافة حاليا.")
        return
    lines = ["المشاريع الحالية:"]
    for p in projects:
        target = float(p["target_amount"])
        collected = float(p["collected"])
        remaining = max(target - collected, 0)
        progress = (collected / target * 100) if target > 0 else 0
        lines.append(
            f"\n#{p['id']} - {p['title']}\n"
            f"المستهدف: {money(p['target_amount'])}\n"
            f"المحصل: {money(p['collected'])}\n"
            f"المتبقي: {money(remaining)}\n"
            f"نسبة التغطية: {progress:.1f}%\n"
            f"التفاصيل: {p['details']}"
        )
    await update.message.reply_text("\n".join(lines))


async def show_income_report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = STORE.approved_transactions()
    totals = STORE.totals()
    if not rows:
        await update.message.reply_text("لا توجد إيرادات معتمدة حتى الآن.")
        return
    lines = [
        "كشف الإيرادات العام",
        f"إجمالي الإيرادات المعتمدة: {money(totals['income'])}",
        "",
        "آخر الإيرادات المعتمدة:",
    ]
    for tx in rows:
        project = f" - {tx['project_title']}" if tx["project_title"] else ""
        reviewed_at = tx["reviewed_at"] or tx["created_at"]
        lines.append(
            f"#{tx['id']} | {tx['payer_name']} | {PAYMENT_TYPES.get(tx['payment_type'])}{project} | "
            f"{money(tx['amount'])} | {reviewed_at}"
        )
    await update.message.reply_text("\n".join(lines))


async def show_expenses(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = STORE.expenses()
    totals = STORE.totals()
    by_category = STORE.expense_totals_by_category()
    if not rows:
        await update.message.reply_text("لا توجد منصرفات مسجلة حتى الآن.")
        return
    lines = [
        "كشف المنصرفات العام",
        f"إجمالي المنصرفات: {money(totals['expenses'])}",
        "",
        "حسب التصنيف:",
    ]
    for row in by_category:
        lines.append(f"- {row['category']}: {money(row['total'])}")
    lines.append("\nآخر المنصرفات:")
    for ex in rows:
        lines.append(
            f"#{ex['id']} | {ex['title']} | {ex['category']} | {money(ex['amount'])} | {ex['created_at']}"
        )
    await update.message.reply_text("\n".join(lines))


async def admin_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    if not is_admin(update.effective_user.id):
        await update.effective_message.reply_text("صفحة الإدارة مخصصة للحسابات المعتمدة فقط.")
        return ConversationHandler.END
    await send_admin_panel(update.effective_message)
    return ConversationHandler.END


async def show_admin_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("صفحة الإدارة مخصصة للحسابات المعتمدة فقط.")
        return
    await send_admin_panel(update.message)


async def send_admin_panel(message):
    totals = STORE.totals()
    await message.reply_text(
        f"لوحة الإدارة\n"
        f"الرصيد: {money(totals['balance'])}\n"
        f"طلبات معلقة: {totals['pending']}",
        reply_markup=admin_keyboard(),
    )


async def admin_actions(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    context.user_data.clear()

    if not is_admin(query.from_user.id):
        await query.edit_message_text("هذه العملية مخصصة للإدارة فقط.")
        return ConversationHandler.END

    action = query.data.split(":", 1)[1]
    if action == "transactions":
        rows = STORE.admin_transactions(12)
        buttons = [[InlineKeyboardButton(f"#{r['id']} {r['payer_name']} · {money(r['amount'])} · {r['status']}", callback_data=f"adm:tx:{r['id']}")] for r in rows]
        buttons.append([InlineKeyboardButton("رجوع", callback_data="admin:home")])
        await query.edit_message_text("آخر المعاملات:", reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if action == "expenses":
        rows = STORE.admin_expenses(12)
        buttons = [[InlineKeyboardButton(f"#{r['id']} {r['title']} · {money(r['amount'])}", callback_data=f"adm:ex:{r['id']}")] for r in rows]
        buttons.append([InlineKeyboardButton("رجوع", callback_data="admin:home")])
        await query.edit_message_text("آخر المنصرفات:", reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if action == "projects":
        rows = STORE.projects(active_only=False)
        buttons = [[InlineKeyboardButton(f"#{r['id']} {r['title']} · {'نشط' if r['active'] else 'متوقف'}", callback_data=f"adm:pr:{r['id']}")] for r in rows]
        buttons += [[InlineKeyboardButton("إضافة مشروع", callback_data="admin:add_project")], [InlineKeyboardButton("رجوع", callback_data="admin:home")]]
        await query.edit_message_text("المشاريع:", reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if action == "search":
        context.user_data["admin_op"] = "search"
        await query.edit_message_text("أرسل كلمة أو رقم البحث (معاملة، دافع، مستخدم، مبلغ، حالة، مشروع، منصرف أو وصف).")
        return ADMIN_INPUT
    if action == "settings":
        buttons=[[InlineKeyboardButton("تعديل قيمة الاشتراك",callback_data="admin:set_subscription")],[InlineKeyboardButton("تعديل بيانات الدفع",callback_data="admin:set_payment_details")],[InlineKeyboardButton("رجوع",callback_data="admin:home")]]
        await query.edit_message_text(f"قيمة الاشتراك: {STORE.setting('monthly_subscription_amount')}\nبيانات الدفع: {STORE.setting('payment_details')}", reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if action == "backup":
        db_path = os.path.abspath(STORE.path)
        os.makedirs(os.path.join(tempfile.gettempdir(), "association_finance_backups"), exist_ok=True)
        backup_path = os.path.join(tempfile.gettempdir(), "association_finance_backups", f"finance_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
        with closing(STORE.connect()) as source, sqlite3.connect(backup_path) as target:
            source.backup(target)
        STORE.audit(query.from_user.id, "backup", "database", None, None, {"file": os.path.basename(backup_path)}, "on-demand backup")
        with open(backup_path, "rb") as backup_file:
            await query.message.reply_document(backup_file, filename=os.path.basename(backup_path), caption="نسخة SQLite متسقة من قاعدة البيانات.")
        await query.edit_message_text("تم إنشاء النسخة الاحتياطية وإرسالها.", reply_markup=admin_keyboard()); return ConversationHandler.END
    if action == "reset":
        await query.edit_message_text("إعادة الضبط ستمسح المعاملات والمنصرفات والمشاريع الحالية. أنشئ نسخة احتياطية أولاً، ثم أرسل كلمة RESET للتأكيد.")
        context.user_data["admin_op"] = "reset_confirm"
        return ADMIN_INPUT
    if action == "home":
        await query.edit_message_text("لوحة الإدارة", reply_markup=admin_keyboard()); return ConversationHandler.END
    if action == "pending":
        await admin_pending(query)
        return ConversationHandler.END
    if action == "report":
        totals = STORE.totals()
        await query.edit_message_text(
            f"تقرير سريع\n"
            f"الإيرادات: {money(totals['income'])}\n"
            f"المنصرفات: {money(totals['expenses'])}\n"
            f"الرصيد: {money(totals['balance'])}\n"
            f"طلبات معلقة: {totals['pending']}",
            reply_markup=admin_keyboard(),
        )
        return ConversationHandler.END
    if action == "monthly_report":
        await query.edit_message_text(monthly_report_text(), reply_markup=admin_keyboard())
        return ConversationHandler.END
    if action == "export_excel":
        await query.edit_message_text("جاري تجهيز ملف Excel...", reply_markup=admin_keyboard())
        path = create_excel_export()
        with open(path, "rb") as file:
            await query.message.reply_document(
                document=file,
                filename=os.path.basename(path),
                caption="تم تجهيز كشف Excel للمالية.",
            )
        return ConversationHandler.END
    if action == "set_subscription":
        context.user_data["setting_key"] = "monthly_subscription_amount"
        await query.edit_message_text("اكتب قيمة الاشتراك الشهري بالأرقام.\nللإلغاء اكتب /cancel")
        return ADMIN_SETTING_VALUE
    if action == "set_payment_details":
        context.user_data["setting_key"] = "payment_details"
        await query.edit_message_text("اكتب تفاصيل الدفع التي ستظهر للأعضاء.\nللإلغاء اكتب /cancel")
        return ADMIN_SETTING_VALUE
    if action == "add_project":
        await query.edit_message_text("اكتب اسم المشروع.\nللإلغاء اكتب /cancel")
        return ADMIN_PROJECT_TITLE
    if action == "add_expense":
        await query.edit_message_text("اكتب عنوان المنصرف.\nللإلغاء اكتب /cancel")
        return ADMIN_EXPENSE_TITLE
    return ConversationHandler.END


async def admin_record_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    if not is_admin(query.from_user.id):
        await query.edit_message_text("هذه العملية للإدارة فقط."); return ConversationHandler.END
    _, kind, raw_id = query.data.split(":", 2)
    entity_id = int(raw_id)
    if kind == "tx":
        row = STORE.transaction(entity_id)
        if not row: await query.edit_message_text("المعاملة غير موجودة."); return ConversationHandler.END
        buttons = [[InlineKeyboardButton("تعديل", callback_data=f"adm:edit_tx:{entity_id}"), InlineKeyboardButton("إبطال", callback_data=f"adm:void_tx:{entity_id}")], [InlineKeyboardButton("رجوع", callback_data="admin:transactions")]]
        await query.edit_message_text(tx_summary(row), reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if kind == "ex":
        row = next((r for r in STORE.admin_expenses(100) if r["id"] == entity_id), None)
        if not row: await query.edit_message_text("المنصرف غير موجود."); return ConversationHandler.END
        buttons = [[InlineKeyboardButton("تعديل", callback_data=f"adm:edit_ex:{entity_id}"), InlineKeyboardButton("إبطال", callback_data=f"adm:void_ex:{entity_id}")], [InlineKeyboardButton("رجوع", callback_data="admin:expenses")]]
        await query.edit_message_text(f"منصرف #{entity_id}\nالعنوان: {row['title']}\nالتصنيف: {row['category']}\nالمبلغ: {money(row['amount'])}\nالتفاصيل: {row['details']}\nالحالة: {'مبطل' if row['void_reason'] else 'فعال'}", reply_markup=InlineKeyboardMarkup(buttons)); return ConversationHandler.END
    if kind == "pr":
        row=STORE.project(entity_id)
        if not row: await query.edit_message_text("المشروع غير موجود."); return ConversationHandler.END
        collected=STORE.project_collected(entity_id)
        buttons=[[InlineKeyboardButton("تعديل المشروع",callback_data=f"adm:edit_pr:{entity_id}"),InlineKeyboardButton("تفعيل/تعطيل",callback_data=f"adm:toggle_pr:{entity_id}")],[InlineKeyboardButton("رجوع",callback_data="admin:projects")]]
        await query.edit_message_text(f"مشروع #{entity_id}: {row['title']}\nالحالة: {'نشط' if row['active'] else 'متوقف'}\nالمستهدف: {money(row['target_amount'])}\nالمحصل المعتمد: {money(collected)}\n{row['details']}",reply_markup=InlineKeyboardMarkup(buttons));return ConversationHandler.END
    if kind.startswith("void_"):
        entity_type="transaction" if kind=="void_tx" else "expense"
        context.user_data.update(admin_op="void", entity_type=entity_type, entity_id=entity_id)
        await query.edit_message_text(f"لتأكيد إبطال {entity_type} #{entity_id}، أرسل السبب. لن يُحذف السجل.")
        return ADMIN_REASON
    if kind.startswith("edit_"):
        context.user_data.update(admin_op="edit", entity_type={"edit_tx":"transaction","edit_ex":"expense","edit_pr":"project"}[kind], entity_id=entity_id)
        await query.edit_message_text('أرسل اسم الحقل كما هو موضح بين القوسين، ثم | ثم القيمة.\\nالمعاملة: اسم الدافع (payer_name) / المبلغ (amount) / نوع الدفع (payment_type) / معرف المشروع (project_id)\\nالمنصرف: العنوان (title) / التصنيف (category) / المبلغ (amount) / التفاصيل (details)\\nالمشروع: الاسم (title) / المبلغ المستهدف (target_amount) / التفاصيل (details)')
        return ADMIN_INPUT
    if kind == "toggle_pr":
        row=STORE.project(entity_id); active=0 if row["active"] else 1
        STORE.update_project(entity_id,"active",active,query.from_user.id,"project active state changed")
        await query.edit_message_text("تم تحديث حالة المشروع.",reply_markup=admin_keyboard());return ConversationHandler.END
    return ConversationHandler.END


async def admin_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        context.user_data.clear(); await update.message.reply_text("هذه العملية للإدارة فقط."); return ConversationHandler.END
    op=context.user_data.get("admin_op")
    value=update.message.text.strip()
    if op == "search":
        txs, exps=STORE.search(value,8)
        buttons=[[InlineKeyboardButton(f"#{r['id']} {r['payer_name']} · {money(r['amount'])}",callback_data=f"adm:tx:{r['id']}")] for r in txs]
        buttons += [[InlineKeyboardButton(f"مصروف #{r['id']} {r['title']} · {money(r['amount'])}",callback_data=f"adm:ex:{r['id']}")] for r in exps]
        buttons.append([InlineKeyboardButton("رجوع",callback_data="admin:home")])
        await update.message.reply_text("نتائج محدودة (حتى 8 لكل نوع):" if txs or exps else "لا توجد نتائج.",reply_markup=InlineKeyboardMarkup(buttons))
        context.user_data.clear();return ConversationHandler.END
    if op == "reset_confirm":
        if value != "RESET":
            await update.message.reply_text("لم تتم إعادة الضبط. أرسل RESET حرفياً للتأكيد أو /cancel للإلغاء.");return ADMIN_INPUT
        context.user_data["admin_op"]="reset_reason"
        await update.message.reply_text("أرسل سبب إعادة الضبط لتسجيله في سجل التدقيق.");return ADMIN_REASON
    if op == "edit":
        if "|" not in value:
            await update.message.reply_text("الصيغة المطلوبة: field | value");return ADMIN_INPUT
        field, raw=value.split("|",1); field=field.strip(); raw=raw.strip()
        try:
            if field in {"amount","target_amount"}: raw=amount_from_text(raw)
            elif field=="project_id": raw=int(raw) if raw else None
            elif field=="payment_type" and raw not in PAYMENT_TYPES: raise ValueError
        except ValueError:
            await update.message.reply_text("قيمة غير صالحة لهذا الحقل.");return ADMIN_INPUT
        context.user_data.update(edit_field=field, edit_value=raw)
        await update.message.reply_text("أرسل سبب التعديل للتدقيق.");return ADMIN_REASON
    await update.message.reply_text("انتهت العملية. افتح لوحة الإدارة مجدداً.");context.user_data.clear();return ConversationHandler.END


async def admin_reason(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        context.user_data.clear(); await update.message.reply_text("هذه العملية للإدارة فقط."); return ConversationHandler.END
    reason=update.message.text.strip()
    if len(reason)<3:
        await update.message.reply_text("اكتب سبباً واضحاً (3 أحرف على الأقل).",);return ADMIN_REASON
    op=context.user_data.get("admin_op"); entity=context.user_data.get("entity_type"); entity_id=context.user_data.get("entity_id"); ok=False
    if op=="void":
        ok=STORE.void_transaction(entity_id,update.effective_user.id,reason) if entity=="transaction" else STORE.void_expense(entity_id,update.effective_user.id,reason)
    elif op=="edit":
        field=context.user_data["edit_field"]; val=context.user_data["edit_value"]
        ok=(STORE.edit_transaction(entity_id,field,val,update.effective_user.id,reason) if entity=="transaction" else STORE.edit_expense(entity_id,field,val,update.effective_user.id,reason) if entity=="expense" else STORE.update_project(entity_id,field,val,update.effective_user.id,reason))
    elif op=="reset_reason":
        before=STORE.reset_counts()
        backup_path=os.path.join(tempfile.gettempdir(), f"finance_pre_reset_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
        with closing(STORE.connect()) as source, sqlite3.connect(backup_path) as backup:
            source.backup(backup)
        with closing(STORE.connect()) as conn:
            conn.execute("DELETE FROM transactions");conn.execute("DELETE FROM expenses");conn.execute("DELETE FROM projects");conn.commit()
        STORE.audit(update.effective_user.id,"reset","database",None,before,{"transactions":0,"expenses":0,"projects":0},reason);ok=True
        context.user_data["reset_backup_path"]=backup_path
    elif op=="setting":
        key=context.user_data.get("setting_key"); old=STORE.setting(key); value=context.user_data["setting_value"]
        STORE.set_setting(key,value); STORE.audit(update.effective_user.id,"change_setting","setting",key,{"value":old},{"value":value},reason); ok=True
    reset_backup_path=context.user_data.get("reset_backup_path")
    context.user_data.clear()
    await update.message.reply_text("تم تسجيل العملية في سجل التدقيق." if ok else "تعذر تنفيذ العملية؛ تحقق من السجل والحالة.",reply_markup=MAIN_MENU)
    if reset_backup_path:
        with open(reset_backup_path,"rb") as backup_file:
            await update.message.reply_document(backup_file,filename=os.path.basename(reset_backup_path),caption="نسخة احتياطية أُنشئت تلقائياً قبل إعادة الضبط.")
    await send_admin_panel(update.message);return ConversationHandler.END


def monthly_report_text() -> str:
    month_prefix = datetime.now().strftime("%Y-%m")
    report = STORE.monthly_totals(month_prefix)
    lines = [
        f"تقرير شهر {month_prefix}",
        f"إجمالي الإيرادات: {money(report['income'])}",
        f"عدد الإيرادات المعتمدة: {report['income_count']}",
        f"إجمالي المنصرفات: {money(report['expenses'])}",
        f"عدد المنصرفات: {report['expense_count']}",
        f"صافي الشهر: {money(report['balance'])}",
    ]

    lines.append("\nالمنصرفات حسب التصنيف:")
    if report["by_category"]:
        for row in report["by_category"]:
            lines.append(f"- {row['category']}: {money(row['total'])}")
    else:
        lines.append("- لا توجد منصرفات هذا الشهر.")

    lines.append("\nأعلى التبرعات:")
    if report["by_project"]:
        for row in report["by_project"]:
            lines.append(f"- {row['title']}: {money(row['total'])}")
    else:
        lines.append("- لا توجد تبرعات معتمدة هذا الشهر.")
    return "\n".join(lines)


def create_excel_export() -> str:
    transactions, expenses, projects = STORE.export_data()
    totals = STORE.totals()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    write_sheet(
        summary,
        "ملخص المالية",
        ["البند", "القيمة"],
        [
            ["إجمالي الإيرادات المعتمدة", totals["income"]],
            ["إجمالي المنصرفات", totals["expenses"]],
            ["الرصيد الحالي", totals["balance"]],
            ["طلبات الدفع المعلقة", totals["pending"]],
            ["تاريخ التصدير", now_iso()],
        ],
    )

    income_sheet = workbook.create_sheet("Income")
    write_sheet(
        income_sheet,
        "كشف الإيرادات",
        ["رقم", "الاسم", "النوع", "المشروع", "المبلغ", "الحالة", "ملاحظة الإدارة", "تاريخ الطلب", "تاريخ المراجعة"],
        [
            [
                tx["id"],
                tx["payer_name"],
                PAYMENT_TYPES.get(tx["payment_type"], tx["payment_type"]),
                tx["project_title"] or "",
                tx["amount"],
                tx["status"],
                tx["admin_note"] or "",
                tx["created_at"],
                tx["reviewed_at"] or "",
            ]
            for tx in transactions
        ],
    )

    expense_sheet = workbook.create_sheet("Expenses")
    write_sheet(
        expense_sheet,
        "كشف المنصرفات",
        ["رقم", "العنوان", "التصنيف", "المبلغ", "التفاصيل", "الحالة", "سبب الإبطال", "تاريخ التسجيل"],
        [
            [ex["id"], ex["title"], ex["category"], ex["amount"], ex["details"], "مبطل" if ex["void_reason"] else "فعال", ex["void_reason"] or "", ex["created_at"]]
            for ex in expenses
        ],
    )

    project_sheet = workbook.create_sheet("Projects")
    write_sheet(
        project_sheet,
        "كشف المشاريع",
        ["رقم", "المشروع", "المستهدف", "المحصل", "المتبقي", "التفاصيل", "نشط", "تاريخ الإضافة"],
        [
            [
                p["id"],
                p["title"],
                p["target_amount"],
                p["collected"],
                float(p["target_amount"]) - float(p["collected"]),
                p["details"],
                "نعم" if p["active"] else "لا",
                p["created_at"],
            ]
            for p in projects
        ],
    )

    output_dir = os.path.join(tempfile.gettempdir(), "association_finance_exports")
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"finance_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    workbook.save(path)
    return path


def write_sheet(sheet, title: str, headers: list[str], rows: list[list]):
    sheet.sheet_view.rightToLeft = True
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=4):
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            if isinstance(value, (int, float)) and col != 1:
                cell.number_format = '#,##0.00'

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, len(rows) + 3)}"
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        max_len = len(str(headers[col - 1]))
        for cell in sheet[letter]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 45))
        sheet.column_dimensions[letter].width = max(12, min(max_len + 3, 45))


async def admin_pending(query):
    rows = STORE.pending_transactions()
    if not rows:
        await query.edit_message_text("لا توجد طلبات دفع معلقة.", reply_markup=admin_keyboard())
        return
    await query.edit_message_text("آخر طلبات الدفع المعلقة ستظهر في رسائل منفصلة.", reply_markup=admin_keyboard())
    for tx in rows:
        keyboard = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton("اعتماد", callback_data=f"review:approve:{tx['id']}"),
                    InlineKeyboardButton("رفض", callback_data=f"review:reject:{tx['id']}"),
                ]
            ]
        )
        await query.message.reply_text(tx_summary(tx), reply_markup=keyboard)


async def admin_setting_value(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END
    key=context.user_data.get("setting_key")
    if not key:
        await update.message.reply_text("????? ???? ???????.");return ConversationHandler.END
    value=update.message.text.strip()
    if key=="monthly_subscription_amount":
        try: amount_from_text(value)
        except ValueError:
            await update.message.reply_text("???? ?????? ?????? ??????.");return ADMIN_SETTING_VALUE
    context.user_data.update(admin_op="setting",setting_value=value)
    await update.message.reply_text("???? ??? ????? ??????? ?????? ??????? ?? ???????.")
    return ADMIN_REASON


async def admin_project_title(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END
    title = update.message.text.strip()
    if len(title) < 2:
        await update.message.reply_text("اكتب اسم مشروع واضح.")
        return ADMIN_PROJECT_TITLE
    context.user_data["project_title"] = title
    await update.message.reply_text("اكتب التكلفة/المبلغ المستهدف بالأرقام:")
    return ADMIN_PROJECT_TARGET


async def admin_project_target(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id): return ConversationHandler.END
    try:
        context.user_data["project_target"] = amount_from_text(update.message.text)
    except ValueError:
        await update.message.reply_text("اكتب مبلغ صحيح بالأرقام.")
        return ADMIN_PROJECT_TARGET
    await update.message.reply_text("اكتب تفاصيل المشروع والتكلفة:")
    return ADMIN_PROJECT_DETAILS


async def admin_project_details(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END
    project_id = STORE.create_project_admin(
        context.user_data["project_title"],
        context.user_data["project_target"],
        update.message.text.strip(),
        update.effective_user.id,
    )
    context.user_data.clear()
    await update.message.reply_text(f"تمت إضافة المشروع رقم #{project_id}.", reply_markup=MAIN_MENU)
    await send_admin_panel(update.message)
    return ConversationHandler.END


async def admin_expense_title(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END
    title = update.message.text.strip()
    if len(title) < 2:
        await update.message.reply_text("اكتب عنوان منصرف واضح.")
        return ADMIN_EXPENSE_TITLE
    context.user_data["expense_title"] = title
    await update.message.reply_text("اختار تصنيف المنصرف:", reply_markup=expense_category_keyboard())
    return ADMIN_EXPENSE_CATEGORY


async def admin_expense_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    if not is_admin(query.from_user.id):
        await query.edit_message_text("هذه العملية مخصصة للإدارة فقط.")
        return ConversationHandler.END
    category = query.data.split(":", 1)[1]
    if category not in EXPENSE_CATEGORIES:
        category = "أخرى"
    context.user_data["expense_category"] = category
    await query.edit_message_text(f"التصنيف: {category}\nاكتب مبلغ المنصرف بالأرقام:")
    return ADMIN_EXPENSE_AMOUNT


async def admin_expense_amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id): return ConversationHandler.END
    try:
        context.user_data["expense_amount"] = amount_from_text(update.message.text)
    except ValueError:
        await update.message.reply_text("اكتب مبلغ صحيح بالأرقام.")
        return ADMIN_EXPENSE_AMOUNT
    await update.message.reply_text("اكتب تفاصيل المنصرف:")
    return ADMIN_EXPENSE_DETAILS


async def admin_expense_details(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return ConversationHandler.END
    expense_id = STORE.add_expense(
        context.user_data["expense_title"],
        context.user_data.get("expense_category", "أخرى"),
        context.user_data["expense_amount"],
        update.message.text.strip(),
        update.effective_user.id,
    )
    context.user_data.clear()
    await update.message.reply_text(f"تم تسجيل المنصرف رقم #{expense_id}.", reply_markup=MAIN_MENU)
    await send_admin_panel(update.message)
    return ConversationHandler.END


def build_app() -> Application:
    app = Application.builder().token(os.environ["BOT_TOKEN"]).build()

    review_flow = ConversationHandler(
        entry_points=[CallbackQueryHandler(review_transaction, pattern="^review:")],
        states={
            REVIEW_REJECT_REASON: [MessageHandler(filters.TEXT & ~filters.COMMAND, reject_reason)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    admin_flow = ConversationHandler(
        entry_points=[
            CommandHandler("admin", admin_start),
            CallbackQueryHandler(admin_actions, pattern="^admin:"),
            CallbackQueryHandler(admin_record_action, pattern="^adm:"),
        ],
        states={
            ADMIN_SETTING_VALUE: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_setting_value)],
            ADMIN_PROJECT_TITLE: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_project_title)],
            ADMIN_PROJECT_TARGET: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_project_target)],
            ADMIN_PROJECT_DETAILS: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_project_details)],
            ADMIN_EXPENSE_TITLE: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_expense_title)],
            ADMIN_EXPENSE_CATEGORY: [CallbackQueryHandler(admin_expense_category, pattern="^expense_category:")],
            ADMIN_EXPENSE_AMOUNT: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_expense_amount)],
            ADMIN_EXPENSE_DETAILS: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_expense_details)],
            ADMIN_INPUT: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_input)],
            ADMIN_REASON: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_reason)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    payment_flow = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex("^الدفع$"), payment_start)],
        states={
            PAY_TYPE: [CallbackQueryHandler(pay_type_selected, pattern="^paytype:")],
            PAY_PROJECT: [CallbackQueryHandler(pay_project_selected, pattern="^payproject:")],
            PAY_AMOUNT: [MessageHandler(filters.TEXT & ~filters.COMMAND, pay_amount)],
            PAY_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, pay_name)],
            PAY_RECEIPT: [MessageHandler(filters.PHOTO | filters.Document.ALL, pay_receipt)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    app.add_handler(review_flow)
    app.add_handler(admin_flow)
    app.add_handler(payment_flow)
    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, main_router))
    return app


load_dotenv()
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

DB_PATH = os.getenv("DB_PATH", "finance_bot.db")
ADMINS = parse_admins()
STORE = Store(DB_PATH)


if __name__ == "__main__":
    if not os.getenv("BOT_TOKEN"):
        raise RuntimeError("BOT_TOKEN is missing. Copy .env.example to .env and fill it.")
    if not ADMINS:
        raise RuntimeError("ADMIN_IDS is missing. Add at least one Telegram user ID.")
    build_app().run_polling(allowed_updates=Update.ALL_TYPES)
